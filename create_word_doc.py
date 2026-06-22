from docx import Document


doc = Document()
doc.add_heading("Smart AI Research Assistant - Full Code Walkthrough", 0)


def add_section(title, paragraphs):
    doc.add_heading(title, level=1)
    for para in paragraphs:
        doc.add_paragraph(para)


add_section(
    "Project Structure Overview",
    [
        "The repository is intentionally small so every major concern has its own top-level folder. The `app/` directory contains all FastAPI backend logic, `ui/` hosts the Streamlit web client, and the root holds documentation (`README.md`), dependency pins (`requirements.txt`), manual smoke tests, and tooling artifacts like the Postman collection. Empty `data/` and `models/` folders are placeholders for future persistence but remain empty because the current build keeps everything in memory and downloads pretrained models on demand.",
        "The virtual environment under `venv/` stores all Python dependencies. It is standard practice not to manually edit files inside `venv/`; they are only consumed when running the project.",
    ],
)

add_section(
    "app/main.py - FastAPI Orchestrator",
    [
        "Lines 1-17 import FastAPI primitives (`FastAPI`, `UploadFile`, `File`, `Query`), CORS middleware, and typing helpers, then pull in the domain-specific helpers: `parse_document`, `generate_summary`, and the retrieval utilities (`create_faiss_index`, `retrieve_chunks`, `answer_question_with_justification`).",
        'A FastAPI application is instantiated with the title "Smart AI Research Assistant" and wrapped with permissive `CORSMiddleware`, allowing any origin, headers, or methods. That design choice keeps local testing effortless and lets the Streamlit frontend talk to the API without additional configuration.',
        "Three module-level stores capture runtime state: `DocumentStore` maps `doc_id` to the raw text, `RetrievalStore` maps `doc_id` to the tuple `(chunks, embeddings, faiss_index)` so we do not recompute embeddings for every query, and `doc_counter` provides monotonic identifiers (`doc_1`, `doc_2`, ...). Later in the file, a second dictionary `docs_db` is declared for Challenge mode. Although both hold raw text, `docs_db` is the one used downstream by the `/challenge` endpoints.",
        "POST /upload: the `upload_document` coroutine is the ingestion pipeline. It calls `parse_document`, then `generate_summary`, then `create_faiss_index`. The resulting text, summary, chunks, embeddings, and FAISS index are cached under a new `doc_id`, and the ID plus summary are returned to the client.",
        "GET /ask: expects `doc_id` and `question`. It validates the ID, retrieves `(chunks, embeddings, index)` from `RetrievalStore`, fetches the top three passages via `retrieve_chunks`, and pipes them into `answer_question_with_justification`, which runs the RoBERTa QA model on the best chunk and returns answer/confidence/justification.",
        'Model bootstrap: the module instantiates Hugging Face pipelines (`qa_pipeline` with "deepset/roberta-base-squad2" and `gen_pipeline` with "gpt2") plus the `docs_db` store. Loading once keeps per-request latency low.',
        "POST /challenge: takes `doc_id`, grabs the first ~800 characters from `docs_db`, and prompts GPT-2 to generate three logic/comprehension questions. It cleans numbering and returns up to three strings.",
        "QAInput model: Pydantic schema for `/challenge/evaluate` containing `doc_id`, `questions`, and `answers` lists.",
        "POST /challenge/evaluate: rebuilds a fresh FAISS index from the stored text, retrieves the best chunk for each submitted question, gets the model's answer via the QA pipeline, compares it to the user answer, and returns structured feedback plus justification chunks.",
    ],
)

add_section(
    "app/parsers.py - Document Ingestion",
    [
        "Docstring states that the parser outputs both full text and a paragraph-to-page map.",
        "Imports `io`, typing helpers, and `pdfplumber`. pdfplumber is pure Python, so it works smoothly on Python 3.13.",
        "`_parse_pdf` reads bytes with `io.BytesIO`, iterates over pages, runs `extract_text`, splits into paragraphs on blank lines, stores each paragraph, and records which PDF page it came from in `page_map`.",
        "`_parse_txt` decodes UTF-8 text, splits on blank lines, and maps every paragraph to page 1.",
        "`parse_document` is the public entry point that detects file extension and routes to the appropriate helper. Unsupported formats raise `ValueError` so the API can return a helpful message.",
    ],
)

add_section(
    "app/retriever.py - Chunking, Embeddings, and QA",
    [
        "Loads `SentenceTransformer('all-MiniLM-L6-v2')` and stores the embedding dimension for FAISS. This happens once at import time.",
        "`chunk_text` splits documents on blank lines, falls back to single lines if necessary, and enforces a 500-character cap by slicing long paragraphs.",
        "`create_faiss_index` encodes all chunks with MiniLM (float32 arrays), instantiates `faiss.IndexFlatL2`, adds the embeddings, and returns `(chunks, embeddings, index)`.",
        "`retrieve_chunks` embeds the user question, queries FAISS for the top `k` matches, and returns the corresponding chunk strings.",
        "`answer_question_with_justification` runs the RoBERTa QA pipeline on the best chunk and packages answer, confidence score, and justification metadata. If no chunks are provided, it returns a default empty response.",
    ],
)

add_section(
    "app/summarizer.py - BART Wrapper",
    [
        "Defines a single global summarization pipeline with `facebook/bart-large-cnn`.",
        "`generate_summary` truncates the document to 3,000 characters, then calls the pipeline with `max_length=150`, `min_length=60`, and `do_sample=False` for deterministic abstracts.",
    ],
)

add_section(
    "ui/streamlit_app.py - Frontend Client",
    [
        'Configures Streamlit, sets `BACKEND = "http://127.0.0.1:8000"`, and structures the UI into three sections: upload, Ask Anything, and Challenge Me.',
        "Upload: uses `st.file_uploader` and posts the file to `/upload`, displaying the returned summary and storing `doc_id`.",
        "Ask Anything: captures a question, calls `/ask`, and renders answer, confidence, and justification paragraph.",
        "Challenge Me: triggers `/challenge` to get questions, keeps them in `st.session_state`, collects user answers, posts them to `/challenge/evaluate`, and displays per-question feedback with justification chunks.",
    ],
)

add_section(
    "Tests and Utility Scripts",
    [
        "`test_parsers.py` mocks an UploadFile via `SimpleNamespace` to verify parser output.",
        "`test_retriever.py` builds a tiny FAISS index and prints retrieved chunks for a sample question.",
        "`test_summary.py` runs `generate_summary` on repeated text to confirm the BART pipeline works.",
        "`hello.py` is a deprecated FastAPI stub kept commented out.",
        "`Smart-ai Assistant.postman_collection.json` offers placeholder requests for manual API testing.",
    ],
)

add_section(
    "Supporting Files",
    [
        "`README.md` documents features, tech stack, and local setup instructions.",
        "`requirements.txt` pins FastAPI, Streamlit, pdfplumber, transformers, sentence-transformers, torch, numpy, faiss-cpu, python-multipart, and optional SQLAlchemy.",
        "`data/` and `models/` are empty scaffolding directories reserved for future persistence of documents or model artifacts.",
    ],
)

OUTPUT_PATH = "Smart_AI_Code_Walkthrough.docx"
doc.save(OUTPUT_PATH)
print("Saved", OUTPUT_PATH)
